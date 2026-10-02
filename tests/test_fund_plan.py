"""月資金計劃表線上表單：模型、公式評估器與 HTTP API 的自動化測試。

執行： python3 -m unittest tests/test_fund_plan.py -v
"""
import datetime as dt
import json
import os
import shutil
import sys
import tempfile
import threading
import time
import unittest
import urllib.request
from http.server import ThreadingHTTPServer

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "src"))
sys.path.insert(0, os.path.join(ROOT, "tests"))

import openpyxl  # noqa: E402

from formula_eval import Evaluator, FormulaError, excel_text  # noqa: E402
from fund_plan_model import (EditError, ExternalModified, FundPlanBook,  # noqa: E402
                             parse_date_input, parse_money_input)
from fund_plan_server import AutoSaver, make_handler  # noqa: E402
from make_fund_plan_sample import build  # noqa: E402


class FormulaTests(unittest.TestCase):
    def ev(self, cells):
        return Evaluator(lambda s, r: cells.get(r))

    def test_arithmetic_and_refs(self):
        ev = self.ev({"F3": 100, "G3": 20, "H3": "=F3-G3", "I3": "=H3", "K3": 50, "L2": 1000,
                      "L3": "=L2+K3-I3"})
        self.assertEqual(ev.value(None, "H3"), 80)
        self.assertEqual(ev.value(None, "L3"), 970)

    def test_functions(self):
        ev = self.ev({"A1": 5, "A2": "", "A3": 7, "B1": '=IF(A2="","",A1)', "B2": "=SUM(A1:A3)",
                      "B3": '=IFERROR(1/0,"x")', "B4": "=ROUND(A1*0.05,0)", "B5": '=TEXT(45900,"aaa")',
                      "B6": "=ABS(-3)&\"元\""})
        self.assertEqual(ev.value(None, "B1"), "")
        self.assertEqual(ev.value(None, "B2"), 12)
        self.assertEqual(ev.value(None, "B3"), "x")
        self.assertEqual(ev.value(None, "B4"), 0)
        self.assertEqual(ev.value(None, "B5"), "日")   # 序號 45900 = 2025-08-31，星期日
        self.assertEqual(ev.value(None, "B6"), "3元")

    def test_text_weekday(self):
        d = dt.datetime(2026, 8, 5)  # 星期三
        self.assertEqual(excel_text(d, "aaa"), "三")
        self.assertEqual(excel_text(d, "aaaa"), "星期三")
        self.assertEqual(excel_text(d, "e/m/d"), "115/8/5")
        self.assertEqual(excel_text(d, "yyyy/mm/dd"), "2026/08/05")
        self.assertEqual(excel_text(1234567, "#,##0"), "1,234,567")

    def test_cycle_and_unsupported(self):
        ev = self.ev({"A1": "=A2", "A2": "=A1", "B1": "=FOO(1)"})
        with self.assertRaises(FormulaError):
            ev.value(None, "A1")
        with self.assertRaises(FormulaError):
            ev.value(None, "B1")


class ParseTests(unittest.TestCase):
    def test_dates(self):
        self.assertEqual(parse_date_input("115/09/10"), dt.datetime(2026, 9, 10))
        self.assertEqual(parse_date_input("2026-9-1"), dt.datetime(2026, 9, 1))
        self.assertEqual(parse_date_input("9/10", 2026), dt.datetime(2026, 9, 10))
        self.assertEqual(parse_date_input("20260910"), dt.datetime(2026, 9, 10))
        self.assertEqual(parse_date_input("1150910"), dt.datetime(2026, 9, 10))
        self.assertIsNone(parse_date_input("  "))
        with self.assertRaises(EditError):
            parse_date_input("abc")

    def test_money(self):
        self.assertEqual(parse_money_input("$1,234,567"), 1234567)
        self.assertEqual(parse_money_input("(500)"), -500)
        self.assertEqual(parse_money_input("12.5"), 12.5)
        self.assertIsNone(parse_money_input(""))
        with self.assertRaises(EditError):
            parse_money_input("十萬")


class BookTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.path = os.path.join(self.tmp, "月資金計劃表.xlsx")
        build(self.path)
        self.book = FundPlanBook(self.path, backup_dir=os.path.join(self.tmp, "備份"))
        self.book.load()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_detect_table_and_columns(self):
        info = self.book.sheets["2026(8)"]
        self.assertEqual(info.table_name, "表8")
        self.assertEqual((info.first_row, info.last_row), (2, 46))
        kinds = {c.letter: (c.kind, c.editable) for c in info.columns}
        self.assertEqual(kinds["C"], ("date", True))
        self.assertEqual(kinds["E"], ("text", True))
        self.assertEqual(kinds["F"], ("money", True))
        self.assertEqual(kinds["H"], ("computed", False))
        self.assertEqual(kinds["L"], ("computed", False))
        self.assertEqual(kinds["A"], ("number", True))
        self.assertEqual(self.book.sheet_names()[:3], ["2026(1)", "2026(2)", "2026(3)"])

    def test_payload_evaluates_formulas(self):
        p = self.book.sheet_payload("2026(8)")
        r = p["rows"][0]["cells"]
        self.assertEqual(r["C"]["d"], "115/08/05")
        self.assertEqual(r["H"]["v"], 386250)
        self.assertEqual(r["L"]["v"], 3_000_000 + 7 * 120_000 - 386250)
        self.assertEqual(r["B"]["d"], "三")
        self.assertTrue(p["rows"][-1]["empty"])
        self.assertFalse(p["rows"][0]["empty"])

    def test_edit_recompute_and_reject_formula_columns(self):
        b = self.book
        b.set_cell("2026(8)", 9, "C", "115/08/28")
        b.set_cell("2026(8)", 9, "E", "台企支存")
        b.set_cell("2026(8)", 9, "F", "1,000")
        p = b.sheet_payload("2026(8)")
        row9 = next(r for r in p["rows"] if r["r"] == 9)
        row8 = next(r for r in p["rows"] if r["r"] == 8)
        self.assertEqual(row9["cells"]["H"]["v"], 1000)
        self.assertEqual(row9["cells"]["L"]["v"], row8["cells"]["L"]["v"] - 1000)
        with self.assertRaises(EditError):
            b.set_cell("2026(8)", 9, "L", "5")
        with self.assertRaises(EditError):
            b.set_cell("2026(8)", 999, "F", "5")
        self.assertFalse(b.set_cell("2026(8)", 9, "F", "1000")["changed"])  # 相同值不記錄
        self.assertEqual(len(b.changes), 3)

    def test_save_writes_original_file_and_log(self):
        b = self.book
        b.set_cell("2026(9)", 8, "C", "9/28")
        b.set_cell("2026(9)", 8, "F", "$1,234,567")
        b.set_cell("2026(9)", 8, "M", "測試")
        res = b.save()
        self.assertTrue(os.path.exists(res["backup"]))
        self.assertFalse(b.dirty)
        wb = openpyxl.load_workbook(self.path)
        ws = wb["2026(9)"]
        self.assertEqual(ws["C8"].value, dt.datetime(2026, 9, 28))
        self.assertEqual(ws["C8"].number_format, "[$-800404]e/m/d;@")
        self.assertEqual(ws["F8"].value, 1234567)
        self.assertEqual(ws["M8"].value, "測試")
        self.assertTrue(str(ws["H8"].value).startswith("="))
        self.assertEqual(ws.tables["表9"].ref, "A1:M46")
        self.assertTrue(wb.calculation.fullCalcOnLoad)
        log = list(wb["異動紀錄"].iter_rows(values_only=True))
        self.assertEqual(log[0][0], "時間")
        self.assertEqual([r[4] for r in log[1:]], ["日期", "支出", "備註"])
        self.assertEqual(log[2][6], "1,234,567")
        # 重新載入後內容一致
        b2 = FundPlanBook(self.path, backup_dir=os.path.join(self.tmp, "備份"))
        b2.load()
        self.assertNotIn("異動紀錄", b2.sheets)
        self.assertEqual(b2.sheet_payload("2026(9)")["rows"][6]["cells"]["F"]["v"], 1234567)

    def test_add_row_extends_table_and_copies_formulas(self):
        b = self.book
        n = b.add_row("2026(8)")
        self.assertEqual(n, 47)
        ws = b.wb["2026(8)"]
        self.assertEqual(ws.tables["表8"].ref, "A1:M47")
        self.assertEqual(ws["L47"].value, "=L46+表8[[#This Row],[未入餘額]]-表8[[#This Row],[應付餘額]]")
        self.assertEqual(ws["H47"].value, "=表8[[#This Row],[支出]]-表8[[#This Row],[提前扣款]]")
        b.set_cell("2026(8)", 47, "F", "10")
        p = b.sheet_payload("2026(8)")
        self.assertEqual(p["rows"][-1]["cells"]["H"]["v"], 10)
        b.save()
        wb = openpyxl.load_workbook(self.path)
        self.assertEqual(wb["2026(8)"].tables["表8"].ref, "A1:M47")

    def test_clear_row(self):
        b = self.book
        recs = b.clear_row("2026(8)", 2)
        self.assertGreaterEqual(len(recs), 4)
        self.assertIsNone(b.wb["2026(8)"]["C2"].value)
        self.assertTrue(str(b.wb["2026(8)"]["L2"].value).startswith("="))

    def test_external_modification_blocks_save(self):
        b = self.book
        b.set_cell("2026(8)", 9, "F", "1")
        time.sleep(0.05)
        os.utime(self.path, None)  # 模擬 Excel 另外存檔
        with self.assertRaises(ExternalModified):
            b.save()
        self.assertTrue(b.status()["external_modified"])
        b.save(force=True)
        self.assertFalse(b.dirty)

    def test_sheet_without_table_falls_back_to_header_scan(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "2026(3)"
        ws.append(["日期", "摘要", "支出", "實際應付", "備註"])
        ws.append([dt.datetime(2026, 3, 5), "彰銀支存", 100, "=C2", ""])
        ws.append([dt.datetime(2026, 3, 8), "台企支存", 50, "=C3", ""])
        path = os.path.join(self.tmp, "plain.xlsx")
        wb.save(path)
        b = FundPlanBook(path, backup_dir=os.path.join(self.tmp, "備份"))
        b.load()
        info = b.sheets["2026(3)"]
        self.assertIsNone(info.table_name)
        self.assertEqual((info.first_row, info.last_row), (2, 3))
        self.assertEqual(b.sheet_payload("2026(3)")["rows"][1]["cells"]["D"]["v"], 50)
        n = b.add_row("2026(3)")
        self.assertEqual(b.wb["2026(3)"][f"D{n}"].value, f"=C{n}")


class ApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.path = os.path.join(cls.tmp, "月資金計劃表.xlsx")
        build(cls.path)
        cls.book = FundPlanBook(cls.path, backup_dir=os.path.join(cls.tmp, "備份"))
        cls.book.load()
        cls.saver = AutoSaver(cls.book)
        cls.saver.start()
        cls.srv = ThreadingHTTPServer(("127.0.0.1", 0), make_handler(cls.book, cls.saver))
        cls.port = cls.srv.server_address[1]
        threading.Thread(target=cls.srv.serve_forever, daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        cls.saver.stop.set()
        cls.srv.shutdown()
        cls.srv.server_close()
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def call(self, path, body=None):
        url = f"http://127.0.0.1:{self.port}{path}"
        req = urllib.request.Request(url, data=json.dumps(body).encode() if body is not None else None,
                                     headers={"Content-Type": "application/json"} if body is not None else {})
        try:
            with urllib.request.urlopen(req) as r:
                return r.status, json.loads(r.read())
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read())

    def test_index_and_workbook(self):
        with urllib.request.urlopen(f"http://127.0.0.1:{self.port}/") as r:
            self.assertIn("月資金計劃表", r.read().decode())
        st, j = self.call("/api/workbook")
        self.assertEqual(st, 200)
        self.assertIn("2026(9)", j["sheets"])
        self.assertEqual(j["sync"], "clean")

    def test_edit_autosaves_to_original_file(self):
        st, j = self.call("/api/cell", {"sheet": "2026(10)", "row": 5, "col": "C", "value": "115/10/25"})
        self.assertEqual(st, 200)
        self.assertTrue(j["changed"])
        self.assertEqual(j["display"], "115/10/25")
        self.assertIn("rows", j["sheet_data"])
        st, j = self.call("/api/cell", {"sheet": "2026(10)", "row": 5, "col": "F", "value": "88,000"})
        self.assertEqual(j["sheet_data"]["rows"][3]["cells"]["H"]["v"], 88000)
        # 公式欄拒絕
        st, j = self.call("/api/cell", {"sheet": "2026(10)", "row": 5, "col": "L", "value": "1"})
        self.assertEqual(st, 400)
        # 等背景寫回
        for _ in range(60):
            time.sleep(0.1)
            if self.call("/api/status")[1]["sync"] == "clean" and self.book.last_saved:
                break
        self.assertIsNotNone(self.book.last_saved)
        wb = openpyxl.load_workbook(self.path)
        self.assertEqual(wb["2026(10)"]["C5"].value, dt.datetime(2026, 10, 25))
        self.assertEqual(wb["2026(10)"]["F5"].value, 88000)
        st, j = self.call("/api/log?limit=10")
        self.assertEqual(j["changes"][0]["col"], "F")

    def test_row_add_clear_and_reload(self):
        st, j = self.call("/api/row/add", {"sheet": "2026(11)"})
        self.assertEqual(st, 200)
        self.assertEqual(j["row"], 47)
        st, j = self.call("/api/row/clear", {"sheet": "2026(11)", "row": 2})
        self.assertEqual(st, 200)
        st, j = self.call("/api/save", {})
        self.assertEqual(st, 200)
        st, j = self.call("/api/reload", {})
        self.assertEqual(st, 200)
        self.assertEqual(j["status"]["sync"], "clean")


if __name__ == "__main__":
    unittest.main()
