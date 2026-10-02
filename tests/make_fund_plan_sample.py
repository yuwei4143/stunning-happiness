#!/usr/bin/env python3
"""產生一份與實際「月資金計劃表」同構的範例工作簿，供線上表單開發、測試與示範。

結構比照 ap-close-tw skill 的 schemas.md「C. 月資金計劃表－2026(N) 分頁」：
  Excel 表格（ListObject），A 序｜B 星期｜C 日期｜D 銀行｜E 摘要｜F 支出｜G 提前扣款｜
  H 實際應付(公式)｜I 應付餘額(公式)｜J 收入｜K 未入餘額｜L 目前現金(公式)｜M 備註
  H = 支出 − 提前扣款、I = 實際應付、L = 上一列 L + 未入餘額 − 應付餘額。

用法： python3 tests/make_fund_plan_sample.py [輸出路徑]   預設 data/月資金計劃表_範例.xlsx
"""
import datetime as dt
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo

ROC = "[$-800404]e/m/d;@"
MONEY = '_-"$"* #,##0_-;\\-"$"* #,##0_-;_-"$"* "-"_-;_-@_-'
HEADERS = ["序", "星期", "日期", "銀行", "摘要", "支出", "提前扣款", "實際應付",
           "應付餘額", "收入", "未入餘額", "目前現金", "備註"]
ROWS_PER_MONTH = 45

SAMPLE = {
    8: [
        ("2026-08-05", "彰銀", "彰銀支存", 386_250, 0, 0, 0, "QN7036338～340 三張"),
        ("2026-08-10", "彰銀", "彰銀交換", 0, 0, 1_850_003, 1_850_003, "0805 送票 9 張"),
        ("2026-08-15", "台企", "台企支存", 212_400, 0, 0, 0, ""),
        ("2026-08-18", "彰銀", "彰銀支存", 1_365_383, 0, 0, 0, "既有 11 張，勿只算新票"),
        ("2026-08-20", "彰銀", "彰銀託收", 0, 0, 640_000, 640_000, ""),
        ("2026-08-25", "彰銀", "彰銀支存", 2_180_000, 0, 0, 0, "登記簿合計 2,018,000，待核對"),
        ("2026-08-31", "聯邦", "聯邦支存", 98_700, 0, 0, 0, ""),
    ],
    9: [
        ("2026-09-05", "彰銀", "彰銀支存", 452_900, 0, 0, 0, ""),
        ("2026-09-10", "彰銀", "彰銀交換", 0, 0, 2_310_500, 2_310_500, "0905 送票"),
        ("2026-09-15", "台企", "台企支存", 174_320, 0, 0, 0, ""),
        ("2026-09-20", "彰銀", "彰銀支存", 1_459_109, 0, 0, 0, "建承 T/T 另備"),
        ("2026-09-25", "彰銀", "彰銀支存", 903_250, 0, 0, 0, ""),
        ("2026-09-30", "聯邦", "聯邦支存", 120_000, 0, 0, 0, ""),
    ],
    10: [
        ("2026-10-05", "彰銀", "彰銀支存", 610_000, 0, 0, 0, ""),
        ("2026-10-15", "台企", "台企支存", 205_000, 0, 0, 0, ""),
        ("2026-10-20", "彰銀", "彰銀交換", 0, 0, 1_200_000, 1_200_000, ""),
    ],
}
OPENING = {m: 3_000_000 + (m - 1) * 120_000 for m in range(1, 13)}


def build(path: str) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    hfill = PatternFill("solid", fgColor="1F4E78")
    hfont = Font(name="Microsoft JhengHei", bold=True, color="FFFFFF", size=10)
    bfont = Font(name="Microsoft JhengHei", size=10)
    for month in range(1, 13):
        ws = wb.create_sheet(f"2026({month})")
        ws.append(HEADERS)
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(1, c)
            cell.fill = hfill
            cell.font = hfont
            cell.alignment = Alignment(horizontal="center")
        ws["O1"] = "期初現金"
        ws["P1"] = OPENING[month]
        ws["P1"].number_format = MONEY
        ws["O1"].font = ws["P1"].font = bfont
        data = SAMPLE.get(month, [])
        first, last = 2, 1 + ROWS_PER_MONTH
        t = f"表{month}"
        for i, r in enumerate(range(first, last + 1)):
            ws.cell(r, 1, i + 1)
            ws.cell(r, 2, f'=IF({get_column_letter(3)}{r}="","",TEXT({get_column_letter(3)}{r},"aaa"))')
            if i < len(data):
                d, bank, summ, out, early, income, unin, memo = data[i]
                ws.cell(r, 3, dt.datetime.fromisoformat(d))
                ws.cell(r, 4, bank)
                ws.cell(r, 5, summ)
                ws.cell(r, 6, out or None)
                ws.cell(r, 7, early or None)
                ws.cell(r, 10, income or None)
                ws.cell(r, 11, unin or None)
                ws.cell(r, 13, memo or None)
            ws.cell(r, 8, f"={t}[[#This Row],[支出]]-{t}[[#This Row],[提前扣款]]")
            ws.cell(r, 9, f"={t}[[#This Row],[實際應付]]")
            prev = "$P$1" if r == first else f"L{r - 1}"
            ws.cell(r, 12, f"={prev}+{t}[[#This Row],[未入餘額]]-{t}[[#This Row],[應付餘額]]")
            ws.cell(r, 3).number_format = ROC
            for c in (6, 7, 8, 9, 10, 11, 12):
                ws.cell(r, c).number_format = MONEY
            for c in range(1, len(HEADERS) + 1):
                ws.cell(r, c).font = bfont
        ref = f"A1:{get_column_letter(len(HEADERS))}{last}"
        table = Table(displayName=t, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        cols = []
        for i, h in enumerate(HEADERS, 1):
            tc = TableColumn(id=i, name=h)
            if h == "實際應付":
                tc.calculatedColumnFormula = TableFormula(attr_text=f"{t}[[#This Row],[支出]]-{t}[[#This Row],[提前扣款]]")
            elif h == "應付餘額":
                tc.calculatedColumnFormula = TableFormula(attr_text=f"{t}[[#This Row],[實際應付]]")
            cols.append(tc)
        table.tableColumns = cols
        ws.add_table(table)
        widths = [5, 6, 11, 8, 14, 13, 11, 13, 13, 13, 13, 14, 30]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)
    print("已產生", path)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "data/月資金計劃表_範例.xlsx")
