"""黃金驗證：用三個辨識批次跑預填，逐欄與使用者修正版 11505_updated 比對。

量測「產出符合我修正過的樣子」的達成率，並列出每一處差異，
讓需人工判斷的欄位（金額勾稽、合併票到期取捨等）一目了然。
"""
import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

import openpyxl
from payable_filler import process_batch, KNOWN_BATCHES  # noqa: E402
from normalize import roc_to_date  # noqa: E402

EXPECTED = "data/11505_expected.xlsx"

# 明細表欄位索引（1-based）
COL = {"支票號碼": 11, "支票帳號": 10, "到期日": 12, "備註": 13,
       "收票日": 14, "實收金額": 15, "送票銀行": 16, "送票日": 17,
       "送票帳號": 18, "用途": 19, "合計": 8}

# 這些欄位由固定規則決定，必須 100% 相符。
DETERMINISTIC = ["支票帳號", "收票日", "送票銀行", "送票日", "送票帳號",
                 "實收金額", "用途"]
# 這些欄位含人工判斷成分，僅報告差異、不計入必過。
JUDGMENT = ["到期日", "備註"]


def load_expected_by_check():
    wb = openpyxl.load_workbook(EXPECTED, data_only=True)
    ws = wb["11505"]
    m = {}
    for r in range(3, ws.max_row + 1):
        tk = ws.cell(r, COL["支票號碼"]).value
        if not tk:
            continue
        row = {name: ws.cell(r, idx).value for name, idx in COL.items()}
        parts = str(tk).replace("、", ",").split(",")
        row["_combined"] = len(parts) > 1  # 一列含多張票 = 合併票（人工步驟）
        for t in parts:
            m[t.strip()] = row
    return m


def norm(v):
    d = roc_to_date(v)
    if d is not None:
        return d
    if v is None:
        return None
    return str(v).strip()


def main():
    exp = load_expected_by_check()
    det_ok = det_total = 0
    mismatches = []
    covered = 0
    for path, batch in KNOWN_BATCHES.items():
        for fc in process_batch(path, batch):
            e = exp.get(fc.支票號碼)
            if e is None:
                continue  # 該票不在此樣板（屬其他批次），跳過
            covered += 1
            got = {
                "支票帳號": fc.支票帳號, "到期日": fc.到期日, "備註": fc.備註,
                "收票日": fc.收票日, "實收金額": fc.實收金額,
                "送票銀行": fc.送票銀行, "送票日": fc.送票日,
                "送票帳號": fc.送票帳號, "用途": fc.用途,
            }
            # 合併票的「實收金額/到期日」屬人工合併步驟，不列入固定規則必過。
            judge = set(JUDGMENT)
            det = list(DETERMINISTIC)
            if e.get("_combined"):
                for nm in ("實收金額", "到期日"):
                    if nm in det:
                        det.remove(nm)
                    judge.add(nm)
            for name in det:
                det_total += 1
                if norm(got[name]) == norm(e[name]):
                    det_ok += 1
                else:
                    mismatches.append((fc.支票號碼, name, got[name], e[name],
                                       "DET"))
            for name in judge:
                if norm(got[name]) != norm(e[name]):
                    mismatches.append((fc.支票號碼, name, got[name], e[name],
                                       "JUDGE"))

    print(f"對應到樣板的支票：{covered} 筆")
    print(f"固定規則欄位相符：{det_ok}/{det_total} "
          f"({det_ok/det_total*100:.1f}%)")

    det_mis = [m for m in mismatches if m[4] == "DET"]
    jud_mis = [m for m in mismatches if m[4] == "JUDGE"]

    if det_mis:
        print("\n[!] 固定規則欄位不符（需修規則）：")
        for tk, name, g, e, _ in det_mis:
            print(f"  {tk} {name}: got={g!r} exp={e!r}")
    if jud_mis:
        print("\n[i] 人工判斷欄位差異（預期，交由使用者最後確認）：")
        for tk, name, g, e, _ in jud_mis:
            print(f"  {tk} {name}: got={g!r} exp={e!r}")

    print("\n結果：", "PASS ✅" if not det_mis else "FAIL ❌（固定規則有誤）")
    sys.exit(0 if not det_mis else 1)


if __name__ == "__main__":
    main()
