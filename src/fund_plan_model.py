"""月資金計劃表工作簿模型：讀取 2026(N) 分頁表格、套用線上編輯、即時寫回原檔。

設計原則（沿用帳務系統鐵則）：
  * 沿用原檔名、原路徑寫回，不另存 xxx_修改版；寫入前先在 備份/ 留一份當次工作階段的備份。
  * 公式欄（實際應付／應付餘額／目前現金…）一律唯讀，由 formula_eval 即時算給畫面看，
    存檔時保留原公式字串並要求 Excel 開檔時全量重算。
  * 每一筆變動都寫進工作簿的「異動紀錄」分頁（時間／分頁／列／欄／原值／新值），
    這就是這份檔案的版本歷程。
  * 原檔若在我們載入後被別人（例如 Excel）另外存檔，就拒絕覆寫並回報衝突，不做靜默合併。
"""
from __future__ import annotations

import copy
import datetime as dt
import os
import re
import shutil
import threading
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter, column_index_from_string, range_boundaries

from formula_eval import Evaluator, FormulaError

ROC_DATE_FMT = "[$-800404]e/m/d;@"
MONEY_FMT = "#,##0"
LOG_SHEET = "異動紀錄"
LOG_HEADERS = ["時間", "分頁", "列", "欄", "欄位", "原值", "新值", "來源"]

PLAN_SHEET_RE = re.compile(r"^(\d{4})\s*\(\s*(\d{1,2})\s*\)\s*$")
VERSION_SHEET_RE = re.compile(r"^\d{8}$")
CELL_RE = re.compile(r"^([A-Z]{1,3})(\d+)$")

DATE_WORDS = ("日期", "到期", "兌現")
MONEY_WORDS = ("支出", "扣款", "應付", "餘額", "現金", "金額", "收入", "存入", "入帳", "合計", "小計")


class ExternalModified(Exception):
    """原檔在載入後被外部程式改過，為避免覆寫他人資料而拒絕存檔。"""


class EditError(ValueError):
    """使用者輸入無法套用（唯讀欄、格式錯誤、列不在表格內）。"""


@dataclass
class Column:
    letter: str
    idx: int
    title: str
    kind: str            # date / money / text / computed
    editable: bool
    formula_hint: str = ""


@dataclass
class SheetInfo:
    name: str
    header_row: int
    first_row: int
    last_row: int          # 最後一列資料列（含表格內空白列）
    min_col: int
    max_col: int
    columns: List[Column]
    table_name: Optional[str] = None
    totals_rows: int = 0
    col_by_title: Dict[str, Column] = field(default_factory=dict)


# ---------------------------------------------------------------------------
# 值的顯示與解析
# ---------------------------------------------------------------------------
def to_roc(d: dt.date) -> str:
    return f"{d.year - 1911}/{d.month:02d}/{d.day:02d}"


def parse_date_input(text: str, default_year: Optional[int] = None) -> Optional[dt.datetime]:
    """接受 115/09/10、115.9.10、2026-09-10、2026/9/10、9/10、20260910、1150910。空白→None。"""
    s = (text or "").strip().replace("－", "-").replace("／", "/")
    if not s:
        return None
    m = re.match(r"^(\d{1,4})[-/.](\d{1,2})[-/.](\d{1,2})$", s)
    if m:
        y, mo, d = (int(x) for x in m.groups())
        if y < 1000:
            y += 1911
        return dt.datetime(y, mo, d)
    m = re.match(r"^(\d{1,2})[-/.](\d{1,2})$", s)
    if m:
        y = default_year or dt.date.today().year
        return dt.datetime(y, int(m.group(1)), int(m.group(2)))
    m = re.match(r"^(\d{4})(\d{2})(\d{2})$", s)
    if m:
        return dt.datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.match(r"^(\d{3})(\d{2})(\d{2})$", s)
    if m:
        return dt.datetime(int(m.group(1)) + 1911, int(m.group(2)), int(m.group(3)))
    raise EditError(f"日期格式無法辨識：{text!r}（可用 115/09/10 或 2026-09-10）")


def parse_money_input(text: str) -> Optional[float]:
    s = (text or "").strip().replace(",", "").replace("$", "").replace("NT", "").replace("＄", "")
    s = s.replace("，", "").replace("元", "")
    if not s:
        return None
    if s.startswith("(") and s.endswith(")"):
        s = "-" + s[1:-1]
    try:
        v = float(s)
    except ValueError:
        raise EditError(f"金額格式無法辨識：{text!r}")
    return int(v) if v.is_integer() else v


def display_value(v: Any, kind: str) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (dt.datetime, dt.date)):
        return to_roc(v)
    if isinstance(v, (int, float)):
        if kind == "number":
            return str(int(v)) if float(v).is_integer() else str(v)
        if kind == "date":
            try:
                return to_roc(dt.date(1899, 12, 30) + dt.timedelta(days=int(v)))
            except (OverflowError, ValueError):
                return str(v)
        if float(v).is_integer():
            return f"{int(v):,}"
        return f"{v:,.2f}"
    return str(v)


def _json_safe(v: Any) -> Any:
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _is_date_format(fmt: str) -> bool:
    f = (fmt or "").lower()
    if f == "general":
        return False
    return ("y" in f or "e/" in f or "e/m" in f or "m/d" in f or "d/m" in f or "mm-" in f
            or "yy" in f or "月" in f or "日" in f)


def _guess_kind(title: str) -> str:
    t = (title or "").replace(" ", "")
    if any(w in t for w in DATE_WORDS):
        return "date"
    if any(w in t for w in MONEY_WORDS):
        return "money"
    return "text"


# ---------------------------------------------------------------------------
# 工作簿
# ---------------------------------------------------------------------------
class FundPlanBook:
    def __init__(self, path: str, backup_dir: Optional[str] = None,
                 log_sheet: Optional[str] = LOG_SHEET, actor: str = "線上表單"):
        self.path = os.path.abspath(path)
        self.backup_dir = backup_dir or os.path.join(os.path.dirname(self.path), "備份")
        self.log_sheet_name = log_sheet
        self.actor = actor
        self.lock = threading.RLock()
        self.wb: Optional[openpyxl.Workbook] = None
        self.loaded_mtime: float = 0.0
        self.loaded_at: Optional[dt.datetime] = None
        self.sheets: Dict[str, SheetInfo] = {}
        self.cached: Dict[Tuple[str, str], Any] = {}
        self.dirty = False
        self.pending: int = 0
        self.changes: List[Dict[str, Any]] = []
        self.last_saved: Optional[dt.datetime] = None
        self.last_error: Optional[str] = None
        self.backed_up: Optional[str] = None
        self._evaluators: Dict[str, Evaluator] = {}
        self._sheet_for_row_context: Optional[SheetInfo] = None

    # ------------------------------------------------------------------ 載入
    def load(self) -> None:
        with self.lock:
            keep_vba = self.path.lower().endswith(".xlsm")
            self.wb = openpyxl.load_workbook(self.path, keep_vba=keep_vba)
            st = os.stat(self.path)
            self.loaded_mtime = st.st_mtime
            self.loaded_at = dt.datetime.now()
            self.sheets = {}
            for ws in self.wb.worksheets:
                if ws.title == self.log_sheet_name or VERSION_SHEET_RE.match(ws.title):
                    continue
                info = self._detect(ws)
                if info:
                    self.sheets[ws.title] = info
            self.cached = self._read_cached_values()
            self._evaluators = {}
            self.dirty = False
            self.pending = 0
            self.last_error = None

    def _read_cached_values(self) -> Dict[Tuple[str, str], Any]:
        """用 read_only + data_only 讀取公式欄的快取值，當公式無法評估時的備援顯示。"""
        out: Dict[Tuple[str, str], Any] = {}
        try:
            wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        except Exception:
            return out
        try:
            for name, info in self.sheets.items():
                if name not in wb.sheetnames:
                    continue
                ws = wb[name]
                comp = [c for c in info.columns if c.kind == "computed"]
                if not comp:
                    continue
                for row in ws.iter_rows(min_row=info.first_row, max_row=info.last_row,
                                        min_col=info.min_col, max_col=info.max_col):
                    for cell in row:
                        col, r = getattr(cell, "column", None), getattr(cell, "row", None)
                        if cell.value is not None and col and r:
                            out[(name, f"{get_column_letter(col)}{r}")] = cell.value
        finally:
            wb.close()
        return out

    def _detect(self, ws) -> Optional[SheetInfo]:
        table = None
        if getattr(ws, "tables", None):
            # 取範圍最大的表格
            table = max(ws.tables.values(), key=lambda t: (range_boundaries(t.ref)[3] - range_boundaries(t.ref)[1]))
        if table is not None:
            c0, r0, c1, r1 = range_boundaries(table.ref)
            header_row = r0
            totals = int(table.totalsRowCount or 0)
            first = r0 + int(table.headerRowCount if table.headerRowCount is not None else 1)
            last = r1 - totals
            calc = {}
            for tc in table.tableColumns:
                f = getattr(tc, "calculatedColumnFormula", None)
                if f is not None and getattr(f, "attr_text", None):
                    calc[tc.name] = "=" + f.attr_text
            cols = self._build_columns(ws, header_row, first, last, c0, c1, calc)
            if not cols:
                return None
            info = SheetInfo(ws.title, header_row, first, last, c0, c1, cols, table.name, totals)
        else:
            found = None
            for r in range(1, 12):
                vals = [str(ws.cell(r, c).value or "") for c in range(1, min(ws.max_column, 40) + 1)]
                if any("摘要" in v for v in vals) and any("日期" in v for v in vals):
                    found = r
                    break
            if found is None:
                return None
            header_row = found
            c0, c1 = 1, min(ws.max_column, 40)
            while c1 > 1 and ws.cell(header_row, c1).value in (None, ""):
                c1 -= 1
            while c0 < c1 and ws.cell(header_row, c0).value in (None, ""):
                c0 += 1
            last = header_row
            blank_run = 0
            r = header_row + 1
            while r < header_row + 5000 and blank_run < 30:
                if any(ws.cell(r, c).value not in (None, "") for c in range(c0, c1 + 1)):
                    last = r
                    blank_run = 0
                else:
                    blank_run += 1
                r += 1
            cols = self._build_columns(ws, header_row, header_row + 1, last, c0, c1, {})
            if not cols:
                return None
            info = SheetInfo(ws.title, header_row, header_row + 1, last, c0, c1, cols)
        info.col_by_title = {c.title: c for c in info.columns}
        return info

    def _build_columns(self, ws, header_row, first, last, c0, c1, calc) -> List[Column]:
        cols: List[Column] = []
        for c in range(c0, c1 + 1):
            title = ws.cell(header_row, c).value
            title = str(title).strip() if title not in (None, "") else get_column_letter(c)
            has_formula = title in calc
            hint = calc.get(title, "")
            if not has_formula:
                for r in range(first, min(last, first + 400) + 1):
                    v = ws.cell(r, c).value
                    if isinstance(v, str) and v.startswith("="):
                        has_formula = True
                        hint = v
                        break
            if has_formula:
                cols.append(Column(get_column_letter(c), c, title, "computed", False, hint))
                continue
            kind = _guess_kind(title)
            if kind == "text":
                vals = [ws.cell(r, c).value for r in range(first, min(last, first + 400) + 1)]
                vals = [v for v in vals if v not in (None, "")]
                if vals and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals):
                    kind = "number"
            cols.append(Column(get_column_letter(c), c, title, kind, True))
        return cols

    # ------------------------------------------------------------------ 讀取
    def sheet_names(self) -> List[str]:
        def key(n):
            m = PLAN_SHEET_RE.match(n)
            return (0, int(m.group(1)), int(m.group(2))) if m else (1, 0, 0)
        return sorted(self.sheets, key=key)

    def _evaluator(self, sheet: str) -> Evaluator:
        ev = self._evaluators.get(sheet)
        if ev is None:
            ev = Evaluator(self._get_cell, self._resolve_structured)
            self._evaluators[sheet] = ev
        return ev

    def _get_cell(self, sheet: Optional[str], ref: str) -> Any:
        if sheet is None or sheet not in self.wb.sheetnames:
            raise FormulaError(f"#REF! 找不到分頁 {sheet}")
        return self.wb[sheet][ref].value

    def _resolve_structured(self, sref: str, row: int) -> Tuple[Optional[str], str]:
        """[@支出] / 表1[[#This Row],[支出]] / 表1[支出] → 儲存格或範圍。"""
        m = re.match(r"^([^\[]*)\[(.*)\]$", sref, re.S)
        if not m:
            raise FormulaError(f"結構化參照無法解析 {sref}")
        tname, inner = m.group(1).strip(), m.group(2)
        info = None
        if tname:
            for i in self.sheets.values():
                if i.table_name == tname:
                    info = i
                    break
        if info is None:
            info = self._sheet_for_row_context
        if info is None:
            raise FormulaError(f"找不到表格 {tname or sref}")
        this_row = False
        col_name = None
        parts = [p.strip() for p in re.split(r",(?![^\[]*\])", inner)]
        for p in parts:
            if p.startswith("@"):
                this_row = True
                p = p[1:]
            if p.startswith("[") and p.endswith("]"):
                p = p[1:-1]
            if p.lower() in ("#this row", "#this_row"):
                this_row = True
                continue
            if p.startswith("#"):
                raise FormulaError(f"不支援 {p}")
            if p:
                col_name = p.replace("''", "'")
        if col_name is None:
            raise FormulaError(f"結構化參照缺少欄名 {sref}")
        col = info.col_by_title.get(col_name)
        if col is None:
            raise FormulaError(f"表格無此欄 {col_name}")
        if this_row:
            if row < info.first_row or row > info.last_row:
                raise FormulaError("#VALUE!")
            return info.name, f"{col.letter}{row}"
        return info.name, f"{col.letter}{info.first_row}:{col.letter}{info.last_row}"

    def _computed(self, sheet: str, ref: str, ev: Evaluator) -> Tuple[Any, str]:
        """回傳 (值, 來源)：eval＝即時算出、cache＝檔案快取值、error＝無法評估。"""
        try:
            return ev.value(sheet, ref), "eval"
        except FormulaError as e:
            if (sheet, ref) in self.cached:
                return self.cached[(sheet, ref)], "cache"
            return f"#{str(e).lstrip('#')}", "error"
        except (RecursionError, ValueError, TypeError, OverflowError):
            if (sheet, ref) in self.cached:
                return self.cached[(sheet, ref)], "cache"
            return "#ERR", "error"

    def sheet_payload(self, name: str) -> Dict[str, Any]:
        with self.lock:
            info = self.sheets.get(name)
            if info is None:
                raise KeyError(name)
            ws = self.wb[name]
            ev = self._evaluator(name)
            ev.invalidate()
            self._sheet_for_row_context = info
            year = None
            m = PLAN_SHEET_RE.match(name)
            if m:
                year = int(m.group(1))
            rows = []
            for r in range(info.first_row, info.last_row + 1):
                cells: Dict[str, Any] = {}
                empty = True
                for col in info.columns:
                    ref = f"{col.letter}{r}"
                    raw = ws[ref].value
                    if col.kind == "computed":
                        val, src = self._computed(name, ref, ev)
                        kind = "date" if isinstance(val, (dt.date, dt.datetime)) else \
                            ("money" if isinstance(val, (int, float)) and not isinstance(val, bool) else "text")
                        entry = {"v": _json_safe(val) if src != "error" else None,
                                 "d": display_value(val, kind) if src != "error" else "",
                                 "src": src, "f": raw if isinstance(raw, str) else None}
                        if src == "error":
                            entry["err"] = str(val)
                        cells[col.letter] = entry
                    else:
                        # 「序」這類流水號欄（number）每列都有值，不列入空白列判斷
                        if raw not in (None, "") and col.kind != "number":
                            empty = False
                        cells[col.letter] = {"v": _json_safe(raw), "d": display_value(raw, col.kind)}
                rows.append({"r": r, "empty": empty, "cells": cells})
            return {
                "sheet": name,
                "year": year,
                "month": int(m.group(2)) if m else None,
                "table": info.table_name,
                "header_row": info.header_row,
                "first_row": info.first_row,
                "last_row": info.last_row,
                "columns": [{"letter": c.letter, "title": c.title, "kind": c.kind,
                             "editable": c.editable, "formula": c.formula_hint} for c in info.columns],
                "rows": rows,
            }

    # ------------------------------------------------------------------ 編輯
    def set_cell(self, sheet: str, row: int, letter: str, text: Any, actor: Optional[str] = None) -> Dict[str, Any]:
        with self.lock:
            info = self.sheets.get(sheet)
            if info is None:
                raise EditError(f"找不到分頁 {sheet}")
            letter = letter.upper()
            col = next((c for c in info.columns if c.letter == letter), None)
            if col is None:
                raise EditError(f"欄 {letter} 不在表格內")
            if not col.editable:
                raise EditError(f"「{col.title}」是公式欄，由 Excel 自動計算，不可手動修改")
            if row < info.first_row or row > info.last_row:
                raise EditError(f"第 {row} 列不在表格範圍內（{info.first_row}–{info.last_row}）")
            ws = self.wb[sheet]
            cell = ws.cell(row, col.idx)
            old = cell.value
            if isinstance(old, str) and old.startswith("="):
                raise EditError("這一格是公式，不可覆寫")
            text = "" if text is None else str(text)
            m = PLAN_SHEET_RE.match(sheet)
            year = int(m.group(1)) if m else None
            if col.kind == "date":
                new = parse_date_input(text, year)
            elif col.kind in ("money", "number"):
                new = parse_money_input(text)
            else:
                new = text.strip() or None
            if _same(old, new):
                return {"changed": False, "row": row, "col": letter}
            cell.value = new
            self._apply_format(ws, info, col, cell, new)
            rec = self._log(sheet, row, col, old, new, actor)
            self._touch(sheet)
            return {"changed": True, "row": row, "col": letter, "old": _json_safe(old),
                    "new": _json_safe(new), "display": display_value(new, col.kind), "log": rec}

    def _apply_format(self, ws, info: SheetInfo, col: Column, cell, new) -> None:
        if new is None:
            return
        fmt = cell.number_format or "General"
        if col.kind == "date":
            if not _is_date_format(fmt):
                cell.number_format = self._column_format(ws, info, col, _is_date_format) or ROC_DATE_FMT
        elif col.kind == "money":
            if fmt == "General":
                cell.number_format = self._column_format(
                    ws, info, col, lambda f: f != "General" and not _is_date_format(f)) or MONEY_FMT

    def _column_format(self, ws, info: SheetInfo, col: Column, pred) -> Optional[str]:
        for r in range(info.first_row, min(info.last_row, info.first_row + 200) + 1):
            c = ws.cell(r, col.idx)
            if c.value not in (None, "") and pred(c.number_format or "General"):
                return c.number_format
        return None

    def clear_row(self, sheet: str, row: int, actor: Optional[str] = None) -> List[Dict[str, Any]]:
        with self.lock:
            info = self.sheets.get(sheet)
            if info is None:
                raise EditError(f"找不到分頁 {sheet}")
            if row < info.first_row or row > info.last_row:
                raise EditError("列不在表格範圍內")
            ws = self.wb[sheet]
            out = []
            for col in info.columns:
                if not col.editable:
                    continue
                cell = ws.cell(row, col.idx)
                if cell.value in (None, ""):
                    continue
                old = cell.value
                cell.value = None
                out.append(self._log(sheet, row, col, old, None, actor))
            if out:
                self._touch(sheet)
            return out

    def add_row(self, sheet: str, actor: Optional[str] = None) -> int:
        """在表格尾端增加一列：延伸表格範圍、複製上一列的公式與格式。"""
        with self.lock:
            info = self.sheets.get(sheet)
            if info is None:
                raise EditError(f"找不到分頁 {sheet}")
            if info.totals_rows:
                raise EditError("此表格有合計列，請直接在 Excel 內新增列")
            ws = self.wb[sheet]
            src_row = info.last_row
            new_row = info.last_row + 1
            # 確認新列目前是空的，避免蓋到表格外既有資料
            for c in range(info.min_col, info.max_col + 1):
                if ws.cell(new_row, c).value not in (None, ""):
                    raise EditError(f"第 {new_row} 列已有資料，無法延伸表格")
            for col in info.columns:
                src = ws.cell(src_row, col.idx)
                dst = ws.cell(new_row, col.idx)
                if src.has_style:
                    dst._style = copy.copy(src._style)
                if col.kind == "computed":
                    f = src.value if isinstance(src.value, str) and src.value.startswith("=") else None
                    if f is None:
                        # 上一列剛好空白，往上找最近一個有公式的儲存格
                        for r in range(src_row - 1, info.first_row - 1, -1):
                            v = ws.cell(r, col.idx).value
                            if isinstance(v, str) and v.startswith("="):
                                f = Translator(v, origin=f"{col.letter}{r}").translate_formula(f"{col.letter}{src_row}")
                                break
                    if f is None and col.formula_hint:
                        f = col.formula_hint
                    if f is not None:
                        dst.value = Translator(f, origin=f"{col.letter}{src_row}").translate_formula(
                            f"{col.letter}{new_row}")
            if info.table_name:
                table = ws.tables[info.table_name]
                c0, r0, c1, r1 = range_boundaries(table.ref)
                table.ref = f"{get_column_letter(c0)}{r0}:{get_column_letter(c1)}{r1 + 1}"
                if table.autoFilter is not None:
                    table.autoFilter.ref = table.ref
            info.last_row = new_row
            rec = {"time": dt.datetime.now().isoformat(timespec="seconds"), "sheet": sheet,
                   "row": new_row, "col": "", "title": "（新增列）", "old": "", "new": "", "actor": actor or self.actor}
            self.changes.append(rec)
            self._append_log_sheet(rec)
            self._touch(sheet)
            return new_row

    def _log(self, sheet, row, col: Column, old, new, actor) -> Dict[str, Any]:
        rec = {"time": dt.datetime.now().isoformat(timespec="seconds"), "sheet": sheet, "row": row,
               "col": col.letter, "title": col.title,
               "old": display_value(old, col.kind), "new": display_value(new, col.kind),
               "actor": actor or self.actor}
        self.changes.append(rec)
        self._append_log_sheet(rec)
        return rec

    def _append_log_sheet(self, rec: Dict[str, Any]) -> None:
        if not self.log_sheet_name:
            return
        if self.log_sheet_name in self.wb.sheetnames:
            ws = self.wb[self.log_sheet_name]
        else:
            ws = self.wb.create_sheet(self.log_sheet_name)
            ws.append(LOG_HEADERS)
            for i, w in enumerate([20, 10, 6, 5, 12, 18, 18, 12], 1):
                ws.column_dimensions[get_column_letter(i)].width = w
        ws.append([rec["time"].replace("T", " "), rec["sheet"], rec["row"], rec["col"], rec["title"],
                   rec["old"], rec["new"], rec["actor"]])

    def _touch(self, sheet: str) -> None:
        self.dirty = True
        self.pending += 1
        ev = self._evaluators.get(sheet)
        if ev:
            ev.invalidate()
        for other in self._evaluators.values():
            other.invalidate()

    # ------------------------------------------------------------------ 存檔
    def external_modified(self) -> bool:
        try:
            return abs(os.stat(self.path).st_mtime - self.loaded_mtime) > 1e-6
        except FileNotFoundError:
            return False

    def _backup(self) -> str:
        os.makedirs(self.backup_dir, exist_ok=True)
        stem, ext = os.path.splitext(os.path.basename(self.path))
        stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        dst = os.path.join(self.backup_dir, f"{stem}_{stamp}{ext}")
        shutil.copy2(self.path, dst)
        return dst

    def save(self, force: bool = False) -> Dict[str, Any]:
        with self.lock:
            if self.wb is None:
                raise RuntimeError("尚未載入")
            if not force and self.external_modified():
                self.last_error = "原檔已被其他程式修改，為避免覆寫而未寫入。請按「重新載入」後再編輯。"
                raise ExternalModified(self.last_error)
            if self.backed_up is None and os.path.exists(self.path):
                self.backed_up = self._backup()
            self.wb.calculation.fullCalcOnLoad = True
            stem, ext = os.path.splitext(os.path.basename(self.path))
            tmp = os.path.join(os.path.dirname(self.path), f".{stem}.{os.getpid()}.saving{ext}")
            try:
                self.wb.save(tmp)
                os.replace(tmp, self.path)
            except PermissionError:
                if os.path.exists(tmp):
                    try:
                        os.remove(tmp)
                    except OSError:
                        pass
                self.last_error = "原檔正被 Excel 或其他程式開啟中，無法寫入；關閉該檔後會自動重試。"
                raise
            except Exception as e:
                if os.path.exists(tmp):
                    try:
                        os.remove(tmp)
                    except OSError:
                        pass
                self.last_error = f"寫入失敗：{e}"
                raise
            self.loaded_mtime = os.stat(self.path).st_mtime
            self.last_saved = dt.datetime.now()
            self.dirty = False
            self.pending = 0
            self.last_error = None
            return {"saved_at": self.last_saved.isoformat(timespec="seconds"), "path": self.path,
                    "backup": self.backed_up}

    def status(self) -> Dict[str, Any]:
        with self.lock:
            return {
                "path": self.path,
                "file": os.path.basename(self.path),
                "loaded_at": self.loaded_at.isoformat(timespec="seconds") if self.loaded_at else None,
                "dirty": self.dirty,
                "pending": self.pending,
                "last_saved": self.last_saved.isoformat(timespec="seconds") if self.last_saved else None,
                "error": self.last_error,
                "external_modified": self.external_modified(),
                "backup": self.backed_up,
                "changes": len(self.changes),
                "sheets": self.sheet_names(),
            }


def _same(a: Any, b: Any) -> bool:
    if a in (None, "") and b in (None, ""):
        return True
    if isinstance(a, dt.datetime) and isinstance(b, dt.datetime):
        return a.date() == b.date()
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        return float(a) == float(b)
    return a == b
